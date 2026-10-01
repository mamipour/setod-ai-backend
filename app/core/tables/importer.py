"""
CSV/XLSX import for org tables
================================
Reuses DuckDB for CSV parsing (same approach as tabular.py) and openpyxl
for XLSX.  Column type inference maps DuckDB types to our column type set.

`preview_import`  — parse file, return inferred columns + first rows (no DB write).
`parse_rows`      — full parse, return list of row dicts coerced against schema.
"""
from __future__ import annotations

import csv
import io
import os
import re
import tempfile
from typing import Any

import duckdb

from app.core.tables.schema import coerce_value, slugify, COLUMN_TYPES

MAX_IMPORT_BYTES = 10 * 1024 * 1024   # 10 MB
MAX_IMPORT_ROWS  = 50_000
_SAMPLE_ROWS     = 5


class TableImportError(ValueError):
    """Raised on unrecoverable parse failures."""


# ── DuckDB type → our column type ─────────────────────────────────────────────

_DUCK_TO_OURS: dict[str, str] = {
    "VARCHAR":   "text",
    "BOOLEAN":   "checkbox",
    "BIGINT":    "number",
    "HUGEINT":   "number",
    "INTEGER":   "number",
    "SMALLINT":  "number",
    "TINYINT":   "number",
    "UBIGINT":   "number",
    "UINTEGER":  "number",
    "FLOAT":     "number",
    "DOUBLE":    "number",
    "DECIMAL":   "number",
    "DATE":      "date",
    "TIMESTAMP": "datetime",
    "TIMESTAMP WITH TIME ZONE": "datetime",
}


def _duck_type_to_ours(duck_type: str) -> str:
    upper = duck_type.upper().split("(")[0].strip()
    return _DUCK_TO_OURS.get(upper, "text")


_IDENT_BAD = re.compile(r"[^a-z0-9]+")


def _safe_key(raw: str, idx: int) -> str:
    key = _IDENT_BAD.sub("_", raw.strip().lower()).strip("_") or f"col_{idx + 1}"
    if key[0].isdigit():
        key = "c_" + key
    return key[:40]


# ── Core parse helpers ─────────────────────────────────────────────────────────

def _parse_csv_bytes(raw: bytes) -> tuple[list[dict], list[str], list[str]]:
    """Return (inferred_columns, original_headers, duck_types) from raw CSV bytes."""
    with tempfile.TemporaryDirectory() as tmp:
        path = os.path.join(tmp, "data.csv")
        with open(path, "wb") as fh:
            fh.write(raw)

        con = duckdb.connect()
        try:
            esc = path.replace("'", "''")
            try:
                con.execute(
                    f"CREATE TABLE src AS SELECT * FROM read_csv('{esc}', "
                    "sample_size=-1, header=true)"
                )
            except duckdb.Error:
                con.execute(
                    f"CREATE TABLE src AS SELECT * FROM read_csv('{esc}', "
                    "sample_size=-1, header=true, all_varchar=true)"
                )

            described = con.execute("DESCRIBE src").fetchall()
            if not described:
                raise TableImportError("File has no columns")

            original_headers = [row[0] for row in described]
            duck_types       = [row[1] for row in described]
        finally:
            con.close()

    keys = _dedupe_keys([_safe_key(h, i) for i, h in enumerate(original_headers)])
    inferred = []
    for i, (h, dt) in enumerate(zip(original_headers, duck_types)):
        inferred.append({
            "key":  keys[i],
            "name": h.strip() or f"Column {i+1}",
            "type": _duck_type_to_ours(dt),
        })
    return inferred, original_headers, duck_types


def _parse_xlsx_bytes(raw: bytes) -> tuple[list[dict], list[str]]:
    """Return (inferred_columns, original_headers) from XLSX bytes."""
    try:
        import openpyxl
    except ImportError:
        raise TableImportError("openpyxl is required for XLSX import")

    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)

    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise TableImportError("XLSX file is empty")

    original_headers = [str(h).strip() if h is not None else f"Column {i+1}"
                        for i, h in enumerate(header_row)]
    keys = _dedupe_keys([_safe_key(h, i) for i, h in enumerate(original_headers)])
    inferred = []
    for i, h in enumerate(original_headers):
        inferred.append({"key": keys[i], "name": h, "type": "text"})  # XLSX: always text on infer

    return inferred, original_headers


def _dedupe_keys(keys: list[str]) -> list[str]:
    seen: set[str] = set()
    out = []
    for k in keys:
        orig, n = k, 2
        while k in seen:
            k = f"{orig}_{n}"
            n += 1
        seen.add(k)
        out.append(k)
    return out


# ── Public API ─────────────────────────────────────────────────────────────────

def preview_import(
    raw: bytes,
    *,
    filename: str,
    columns: list[dict],
) -> dict:
    """Parse the file and return inferred column mapping + sample rows.

    If `columns` is provided (table already has schema), shows which headers
    match existing column keys.
    """
    if len(raw) > MAX_IMPORT_BYTES:
        raise TableImportError(f"File exceeds {MAX_IMPORT_BYTES // 1024 // 1024} MB limit")

    ext = os.path.splitext(filename.lower())[-1]
    if ext in (".xlsx", ".xls"):
        inferred, orig_headers = _parse_xlsx_bytes(raw)
    else:
        inferred, orig_headers, _ = _parse_csv_bytes(raw)

    existing_keys = {c["key"] for c in columns}
    for col in inferred:
        col["matches_existing"] = col["key"] in existing_keys

    # Read sample rows
    sample = _read_sample_rows(raw, filename, orig_headers, n=_SAMPLE_ROWS)

    return {
        "inferred_columns": inferred,
        "sample_rows":      sample,
        "existing_columns": columns,
    }


def parse_rows(
    raw: bytes,
    *,
    filename: str,
    columns: list[dict],
) -> list[dict]:
    """Full parse: return a list of row dicts coerced against the table schema.

    Unknown columns are dropped.  Type mismatches fall back leniently to text
    (the schema enforces correctness at create_row time).  Rows beyond MAX_IMPORT_ROWS
    are silently truncated.
    """
    if len(raw) > MAX_IMPORT_BYTES:
        raise TableImportError(f"File exceeds {MAX_IMPORT_BYTES // 1024 // 1024} MB limit")

    ext = os.path.splitext(filename.lower())[-1]
    if ext in (".xlsx", ".xls"):
        return _parse_xlsx_rows(raw, columns)
    else:
        return _parse_csv_rows(raw, columns)


# ── Internal row parsers ───────────────────────────────────────────────────────

def _parse_csv_rows(raw: bytes, columns: list[dict]) -> list[dict]:
    col_map = {c["key"]: c for c in columns}

    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")

    reader = csv.DictReader(io.StringIO(text))
    # Build key mapping: header name → our column key (if header matches a key or name)
    name_to_key = {}
    for c in columns:
        name_to_key[c["key"]]  = c["key"]
        name_to_key[c["name"]] = c["key"]

    rows = []
    for raw_row in reader:
        row_data: dict[str, Any] = {}
        for header, raw_val in raw_row.items():
            if header is None:
                continue
            key = name_to_key.get(header.strip()) or name_to_key.get(_safe_key(header.strip(), 0))
            if key and key in col_map:
                try:
                    row_data[key] = coerce_value(key, col_map[key], raw_val, lenient=True)
                except Exception:
                    row_data[key] = str(raw_val)[:10_000] if raw_val else None
        if row_data:
            rows.append(row_data)
        if len(rows) >= MAX_IMPORT_ROWS:
            break

    return rows


def _parse_xlsx_rows(raw: bytes, columns: list[dict]) -> list[dict]:
    try:
        import openpyxl
    except ImportError:
        raise ImportError("openpyxl is required for XLSX import")

    col_map = {c["key"]: c for c in columns}
    name_to_key: dict[str, str] = {}
    for c in columns:
        name_to_key[c["key"]]  = c["key"]
        name_to_key[c["name"]] = c["key"]

    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)

    try:
        header_row = next(rows_iter)
    except StopIteration:
        return []

    headers = [str(h).strip() if h is not None else "" for h in header_row]
    key_map = []  # [(col_idx, key, col_def), ...]
    for i, h in enumerate(headers):
        key = name_to_key.get(h) or name_to_key.get(_safe_key(h, i))
        if key and key in col_map:
            key_map.append((i, key, col_map[key]))

    rows = []
    for raw_row in rows_iter:
        row_data: dict[str, Any] = {}
        for col_idx, key, col_def in key_map:
            if col_idx < len(raw_row):
                raw_val = raw_row[col_idx]
                try:
                    row_data[key] = coerce_value(key, col_def, raw_val, lenient=True)
                except Exception:
                    row_data[key] = str(raw_val)[:10_000] if raw_val is not None else None
        if row_data:
            rows.append(row_data)
        if len(rows) >= MAX_IMPORT_ROWS:
            break

    return rows


def _read_sample_rows(raw: bytes, filename: str, orig_headers: list[str], n: int) -> list[dict]:
    """Return up to n rows as dicts keyed by original header."""
    ext = os.path.splitext(filename.lower())[-1]
    try:
        if ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            next(rows_iter)  # skip header
            result = []
            for row in rows_iter:
                result.append({h: (str(v) if v is not None else "") for h, v in zip(orig_headers, row)})
                if len(result) >= n:
                    break
            return result
        else:
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("latin-1")
            reader = csv.DictReader(io.StringIO(text))
            result = []
            for row in reader:
                result.append({h: str(v or "") for h, v in row.items()})
                if len(result) >= n:
                    break
            return result
    except Exception:
        return []
