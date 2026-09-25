"""
Tabular data
============
CSV and XLSX knowledge files become SQL tables the agent can query with one tool,
`query_data(sql)`. The model is far better at writing `SELECT … WHERE … GROUP BY` than at
reading five thousand rows of context and filtering them in its head — and the SQL is
deterministic where the reading is not.

Two halves, split across the trusted/untrusted line:

- **Ingest** (upload request, trusted input from the owner): parse the file with DuckDB
  (CSV) or openpyxl → CSV → DuckDB (XLSX), normalise column names, capture the schema and a
  few sample rows, and write Parquet. External access is *on* here because we are reading
  a temp file we wrote ourselves.
- **Query** (run time, SQL written by the model, possibly under prompt injection): a fresh
  in-memory connection per call, tables loaded from Parquet, then
  `enable_external_access=false` + `lock_configuration=true` so the statement cannot read
  the filesystem, write files, install extensions, or flip those settings back. One
  statement, a memory cap, a timeout enforced with `interrupt()`, and a row/byte cap on the
  result. No Python is ever executed; that is why this needs no sandbox.

The schema goes into the tool *description* at run start. Without it the model guesses
column names and types and the first call is always wasted; with it — plus a sample row so
it can see that `category` holds "IT" rather than "Information Technology" — the first
query is usually right.
"""

from __future__ import annotations

import asyncio
import csv
import io
import logging
import os
import re
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any
from uuid import UUID

import duckdb
from sqlmodel import select
from sqlmodel.ext.asyncio.session import AsyncSession

from app.db.models import AgentDataTable, AgentKnowledgeFile

log = logging.getLogger(__name__)

TABULAR_EXTENSIONS = (".csv", ".xlsx")

# Ingest caps. The upload itself is capped at 10 MB upstream; these keep one pathological
# workbook from producing hundreds of tables or a table the description cannot show.
MAX_SHEETS = 10
MAX_COLUMNS = 200
MAX_PARQUET_BYTES = 20 * 1024 * 1024
SAMPLE_ROWS = 3
HEADER_SCAN_ROWS = 10

# Query caps. 256 MB is a quarter of the VPS; a 10 MB CSV materialises to a few tens of MB.
QUERY_MEMORY_LIMIT = "256MB"
QUERY_TIMEOUT_SECONDS = 20
MAX_RESULT_ROWS = 200
MAX_RESULT_CHARS = 32_000
MAX_CELL_CHARS = 200

# Description caps: the schema is prompt text on every run, so it has to stay bounded.
DESCRIPTION_MAX_COLUMNS = 40
DESCRIPTION_MAX_CHARS = 8_000

CACHE_DIR = os.path.join(tempfile.gettempdir(), "setod-tables")


class TabularError(RuntimeError):
    """Surfaced to the uploader (HTTP 422) with a user-readable message."""


@dataclass
class TableDraft:
    """A parsed table before it has a database row."""

    name: str
    sheet: str | None
    row_count: int
    columns: list[dict[str, Any]]  # {"name", "type", "original"}
    sample: list[list[str]]
    parquet: bytes
    # Plain text rendering (header + rows) for the existing chunk/embed pipeline.
    text: str = field(default="", repr=False)


# ── Identifiers ────────────────────────────────────────────────────────────────

_IDENT_BAD = re.compile(r"[^a-z0-9]+")


def _reserved_keywords() -> set[str]:
    con = duckdb.connect()
    try:
        rows = con.execute(
            "SELECT keyword_name FROM duckdb_keywords() WHERE keyword_category = 'reserved'"
        ).fetchall()
    finally:
        con.close()
    return {r[0].lower() for r in rows}


_RESERVED: set[str] | None = None


def sanitise_identifier(raw: str, *, fallback: str) -> str:
    """Lower snake_case, no leading digit, never a reserved word, never empty."""
    global _RESERVED
    if _RESERVED is None:
        _RESERVED = _reserved_keywords()
    ident = _IDENT_BAD.sub("_", raw.strip().lower()).strip("_")
    if not ident:
        ident = fallback
    if ident[0].isdigit():
        ident = f"c_{ident}"
    if ident in _RESERVED:
        ident = f"{ident}_t"
    return ident[:63]


def _dedupe(names: list[str], taken: set[str] | None = None) -> list[str]:
    seen: set[str] = set(taken or ())
    out: list[str] = []
    for n in names:
        candidate, i = n, 2
        while candidate in seen:
            candidate = f"{n}_{i}"
            i += 1
        seen.add(candidate)
        out.append(candidate)
    return out


def table_name_for(filename: str, sheet: str | None, taken: set[str]) -> str:
    stem = os.path.splitext(os.path.basename(filename))[0]
    base = sanitise_identifier(stem, fallback="data")
    if sheet is not None:
        base = f"{base}__{sanitise_identifier(sheet, fallback='sheet')}"
    return _dedupe([base[:63]], taken)[0]


# ── Ingest ─────────────────────────────────────────────────────────────────────

def _read_csv_into(con: duckdb.DuckDBPyConnection, path: str) -> None:
    """Load a CSV as `src`, typed when DuckDB can infer types cleanly, all-text otherwise.

    Real spreadsheets mix `2026-10-03` and `03/11/2026` in one column; strict inference
    then fails on row 2. Text is the honest fallback — the model can `TRY_CAST` or
    `strptime` — and beats refusing the file.
    """
    esc = path.replace("'", "''")
    try:
        con.execute(
            f"CREATE OR REPLACE TABLE src AS SELECT * FROM read_csv('{esc}', "
            "sample_size=-1, header=true)"
        )
    except duckdb.Error:
        con.execute(
            f"CREATE OR REPLACE TABLE src AS SELECT * FROM read_csv('{esc}', "
            "sample_size=-1, header=true, all_varchar=true)"
        )


def _draft_from_csv_path(path: str, *, name: str, sheet: str | None) -> TableDraft:
    con = duckdb.connect()
    try:
        _read_csv_into(con, path)
        described = con.execute("DESCRIBE src").fetchall()
        if not described:
            raise TabularError("The table has no columns.")
        if len(described) > MAX_COLUMNS:
            raise TabularError(f"Tables are limited to {MAX_COLUMNS} columns.")

        originals = [row[0] for row in described]
        types = [row[1] for row in described]
        names = _dedupe([
            sanitise_identifier(o, fallback=f"col_{i + 1}") for i, o in enumerate(originals)
        ])
        select_list = ", ".join(
            f'"{o.replace(chr(34), chr(34) * 2)}" AS "{n}"' for o, n in zip(originals, names)
        )
        con.execute(f"CREATE TABLE t AS SELECT {select_list} FROM src")

        row_count = con.execute("SELECT count(*) FROM t").fetchone()[0]
        if row_count == 0:
            raise TabularError("The table has a header but no rows.")

        sample_rows = con.execute(
            "SELECT " + ", ".join(f'CAST("{n}" AS VARCHAR)' for n in names)
            + f" FROM t LIMIT {SAMPLE_ROWS}"
        ).fetchall()
        sample = [["" if v is None else str(v)[:MAX_CELL_CHARS] for v in r] for r in sample_rows]

        with tempfile.TemporaryDirectory() as tmp:
            out = os.path.join(tmp, "t.parquet")
            con.execute(f"COPY t TO '{out}' (FORMAT PARQUET, COMPRESSION ZSTD)")
            with open(out, "rb") as fh:
                parquet = fh.read()
        if len(parquet) > MAX_PARQUET_BYTES:
            raise TabularError("This table is too large to query.")

        # Text rendering for the embed pipeline: header + rows as CSV, bounded upstream.
        buf = io.StringIO()
        writer = csv.writer(buf)
        writer.writerow(names)
        for r in con.execute("SELECT * FROM t").fetchall():
            writer.writerow(["" if v is None else v for v in r])

        columns = [{"name": n, "type": t, "original": o} for n, t, o in zip(names, types, originals)]
        return TableDraft(
            name=name, sheet=sheet, row_count=row_count, columns=columns,
            sample=sample, parquet=parquet, text=buf.getvalue(),
        )
    finally:
        con.close()


def _detect_header_row(rows: list[list[Any]]) -> int:
    """Index of the first row that looks like a header, within the first few rows.

    Government and supplier workbooks routinely start with a title, a logo row and a blank
    line. A header is the first row where at least half the cells of the widest row are
    filled and every filled cell is text. Falls back to 0.
    """
    width = max((len(r) for r in rows[:HEADER_SCAN_ROWS]), default=0)
    if width == 0:
        return 0
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        filled = [c for c in row if c not in (None, "")]
        if len(filled) >= max(1, (width + 1) // 2) and all(isinstance(c, str) for c in filled):
            return i
    return 0


def _cell_to_csv(value: Any) -> Any:
    if isinstance(value, datetime):
        # Midnight timestamps are how Excel stores plain dates.
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return "" if value is None else value


def _drafts_from_xlsx(data: bytes, filename: str, taken: set[str]) -> list[TableDraft]:
    try:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises a zoo of types
        raise TabularError(f"Could not read this workbook: {exc}") from exc

    drafts: list[TableDraft] = []
    used = set(taken)
    try:
        sheets = wb.worksheets[:MAX_SHEETS]
        for ws in sheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            # Drop fully empty rows; they carry nothing and confuse header detection.
            rows = [r for r in rows if any(c not in (None, "") for c in r)]
            if len(rows) < 2:
                continue
            header_idx = _detect_header_row(rows)
            header = rows[header_idx]
            body = rows[header_idx + 1:]
            if not body:
                continue
            width = max(len(header), max(len(r) for r in body))
            header = [
                (str(h) if h not in (None, "") else f"col_{i + 1}")
                for i, h in enumerate(list(header) + [None] * (width - len(header)))
            ]
            with tempfile.TemporaryDirectory() as tmp:
                path = os.path.join(tmp, "sheet.csv")
                with open(path, "w", newline="", encoding="utf-8") as fh:
                    writer = csv.writer(fh)
                    writer.writerow(header)
                    for r in body:
                        r = list(r) + [None] * (width - len(r))
                        writer.writerow([_cell_to_csv(c) for c in r])
                sheet_name = ws.title if len(sheets) > 1 else None
                name = table_name_for(filename, sheet_name, used)
                used.add(name)
                drafts.append(_draft_from_csv_path(path, name=name, sheet=ws.title))
    finally:
        wb.close()

    if not drafts:
        raise TabularError("The workbook has no sheet with a header row and data.")
    return drafts


def ingest(filename: str, data: bytes, *, taken: set[str]) -> list[TableDraft]:
    """Parse an upload into tables. `taken` is the agent's existing table names.

    Raises TabularError with a message fit for the uploader.
    """
    lower = filename.lower()
    if lower.endswith(".xlsx"):
        return _drafts_from_xlsx(data, filename, taken)
    if lower.endswith(".csv"):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "upload.csv")
            with open(path, "wb") as fh:
                fh.write(data)
            name = table_name_for(filename, None, taken)
            try:
                return [_draft_from_csv_path(path, name=name, sheet=None)]
            except duckdb.Error as exc:
                raise TabularError(f"Could not parse this CSV: {str(exc).splitlines()[0]}") from exc
    raise TabularError("Not a tabular file.")


# ── Description ────────────────────────────────────────────────────────────────

def _fmt_count(n: int) -> str:
    return f"{n:,}"


def render_schema(tables: list[TableMeta], filenames: dict[UUID, str]) -> str:
    """The schema block for the tool description: name, source, columns, one sample row."""
    lines: list[str] = []
    for t in tables:
        source = filenames.get(t.file_id, "upload")
        if t.sheet:
            source = f"{source} › {t.sheet}"
        lines.append(f"- {t.name}  ({source}, {_fmt_count(t.row_count)} rows)")
        cols = t.columns[:DESCRIPTION_MAX_COLUMNS]
        col_text = ", ".join(
            f"{c['name']} {c['type']}" + (f" [was: {c['original']}]" if c["original"] != c["name"] else "")
            for c in cols
        )
        if len(t.columns) > DESCRIPTION_MAX_COLUMNS:
            col_text += f", … +{len(t.columns) - DESCRIPTION_MAX_COLUMNS} more (run DESCRIBE {t.name})"
        lines.append(f"    columns: {col_text}")
        if t.sample:
            row = " | ".join(v[:40] for v in t.sample[0][:DESCRIPTION_MAX_COLUMNS])
            lines.append(f"    e.g. {row}")
    text = "\n".join(lines)
    if len(text) > DESCRIPTION_MAX_CHARS:
        text = text[:DESCRIPTION_MAX_CHARS] + "\n… (schema truncated; run DESCRIBE <table>)"
    return text


TOOL_INTRO = (
    "Run ONE read-only SQL statement (DuckDB dialect, PostgreSQL-like) over the tables "
    "below; the result comes back as a table. Use it to filter, count, join, sort and "
    "aggregate instead of reading whole files. Column types are shown — a date stored as "
    "VARCHAR needs TRY_CAST(col AS DATE) or strptime(col, '%d/%m/%Y'). Column names were "
    "normalised to snake_case; the original header is shown as [was: …]. To explore, use "
    "`SUMMARIZE <table>` (per-column min/max/nulls/distinct) or `SELECT DISTINCT col`. "
    f"Results are capped at {MAX_RESULT_ROWS} rows — aggregate or add LIMIT. An error is "
    "returned verbatim: fix the SQL and call again.\n\nTables:\n"
)


# ── Query ──────────────────────────────────────────────────────────────────────

@dataclass
class TableMeta:
    """Everything about a table except its bytes — what a run needs before any query."""

    id: UUID
    file_id: UUID
    name: str
    sheet: str | None
    row_count: int
    columns: list[dict[str, Any]]
    sample: list[list[str]]


def _cache_file(table_id: UUID) -> str:
    return os.path.join(CACHE_DIR, f"{table_id}.parquet")


def _write_cache(table_id: UUID, parquet: bytes) -> str:
    """Parquet on local disk, keyed by table id; written once, read by every run."""
    os.makedirs(CACHE_DIR, exist_ok=True)
    path = _cache_file(table_id)
    tmp = f"{path}.{os.getpid()}.tmp"
    with open(tmp, "wb") as fh:
        fh.write(parquet)
    os.replace(tmp, path)  # atomic: a concurrent run never sees a half-written file
    return path


def _format_result(columns: list[str], rows: list[tuple], *, truncated: bool) -> str:
    def cell(v: Any) -> str:
        s = "" if v is None else str(v)
        s = s.replace("\n", " ").replace("|", "\\|")
        return s if len(s) <= MAX_CELL_CHARS else s[: MAX_CELL_CHARS - 1] + "…"

    lines = ["| " + " | ".join(columns) + " |", "|" + "---|" * len(columns)]
    for r in rows:
        lines.append("| " + " | ".join(cell(v) for v in r) + " |")
    out = "\n".join(lines)
    if len(out) > MAX_RESULT_CHARS:
        out = out[:MAX_RESULT_CHARS] + "\n… (output truncated)"
    if truncated:
        out += (
            f"\n\nShowing the first {MAX_RESULT_ROWS} rows of a longer result. "
            "Aggregate (COUNT, GROUP BY) or add a tighter WHERE / LIMIT."
        )
    return out + f"\n\n({len(rows)} row{'s' if len(rows) != 1 else ''})"


def run_query(
    tables: list[tuple[str, str]], sql: str, *, handle: dict[str, Any] | None = None
) -> str:
    """Execute one statement against Parquet-backed tables. Synchronous; see `query`.

    `tables` is (name, parquet_path). Returns the model-facing text — results or a readable
    error. Never raises for anything the SQL did; only for our own bugs. When `handle` is
    given, the live connection is published under `handle["con"]` so another thread can
    `interrupt()` it.
    """
    sql = sql.strip().rstrip(";").strip()
    if not sql:
        return "Provide a SQL statement."

    con = duckdb.connect(config={"memory_limit": QUERY_MEMORY_LIMIT, "threads": 2})
    if handle is not None:
        handle["con"] = con
    try:
        for name, path in tables:
            con.execute(f'CREATE TABLE "{name}" AS SELECT * FROM read_parquet(?)', [path])
        # From here on, nothing the statement does can touch the outside world.
        con.execute("SET enable_external_access = false")
        con.execute("SET lock_configuration = true")

        try:
            if len(con.extract_statements(sql)) != 1:
                return "Only one SQL statement per call. Send the statements one at a time."
        except duckdb.Error as exc:
            return f"SQL error: {str(exc).splitlines()[0]}"

        try:
            cur = con.execute(sql)
            if cur.description is None:
                return "The statement ran but returned no rows. Use SELECT to read data."
            columns = [d[0] for d in cur.description]
            rows = cur.fetchmany(MAX_RESULT_ROWS + 1)
        except duckdb.InterruptException:
            return (
                f"The query was stopped after {QUERY_TIMEOUT_SECONDS} seconds. Narrow it with "
                "WHERE, or aggregate instead of selecting every row."
            )
        except duckdb.Error as exc:
            first = str(exc).splitlines()[0]
            return f"SQL error: {first}"

        truncated = len(rows) > MAX_RESULT_ROWS
        return _format_result(columns, rows[:MAX_RESULT_ROWS], truncated=truncated)
    finally:
        con.close()


async def query(tables: list[tuple[str, str]], sql: str) -> str:
    """`run_query` off the event loop, with a wall-clock timeout enforced via interrupt().

    DuckDB is synchronous; running it inline would stall every other agent for the
    duration. The interrupt is what makes the timeout real — cancelling the awaiting task
    alone would leave the query running in the thread.
    """
    handle: dict[str, Any] = {}
    loop = asyncio.get_running_loop()
    fut = loop.run_in_executor(None, lambda: run_query(tables, sql, handle=handle))
    try:
        return await asyncio.wait_for(asyncio.shield(fut), timeout=QUERY_TIMEOUT_SECONDS)
    except asyncio.TimeoutError:
        con = handle.get("con")
        if con is not None:
            try:
                con.interrupt()
            except Exception:  # noqa: BLE001 — connection may already be closed
                pass
        try:
            return await fut  # run_query turns the InterruptException into text
        except Exception:  # noqa: BLE001
            return f"The query was stopped after {QUERY_TIMEOUT_SECONDS} seconds."


# ── Tool ───────────────────────────────────────────────────────────────────────

async def tables_for_agent(db: AsyncSession, agent_id: UUID) -> list[TableMeta]:
    """Table metadata for an agent, deliberately without the Parquet column: the bytes are
    only fetched on a cache miss, so a warm run costs one small query."""
    rows = await db.exec(
        select(
            AgentDataTable.id, AgentDataTable.file_id, AgentDataTable.name, AgentDataTable.sheet,
            AgentDataTable.row_count, AgentDataTable.columns, AgentDataTable.sample,
        )
        .where(AgentDataTable.agent_id == agent_id)
        .order_by(AgentDataTable.created_at)
    )
    return [TableMeta(*r) for r in rows.all()]


async def _ensure_cached(db: AsyncSession, table: TableMeta) -> str:
    path = _cache_file(table.id)
    if os.path.exists(path):
        return path
    row = await db.exec(select(AgentDataTable.parquet).where(AgentDataTable.id == table.id))
    parquet = row.first()
    if parquet is None:
        raise RuntimeError(f"table {table.id} vanished")
    return _write_cache(table.id, parquet)


async def build_tool(db: AsyncSession, agent_id: UUID):
    """The `query_data` tool for one run, or None when the agent has no tables.

    Absent rather than present-but-empty, like `search_knowledge`: a model offered a tool
    will try it.
    """
    from app.core.agents.base import RegisteredTool
    from app.core.llm.client import ToolSpec

    tables = await tables_for_agent(db, agent_id)
    if not tables:
        return None

    file_rows = await db.exec(
        select(AgentKnowledgeFile.id, AgentKnowledgeFile.filename).where(
            AgentKnowledgeFile.id.in_([t.file_id for t in tables])
        )
    )
    filenames = {fid: fname for fid, fname in file_rows.all()}
    # Materialise Parquet to the local cache now, off the hot path of the first query.
    paths = [(t.name, await _ensure_cached(db, t)) for t in tables]

    async def handler(args: dict, dry_run: bool) -> str:
        # Read-only by construction, so dry_run needs no special casing.
        sql = str(args.get("sql", ""))
        return await query(paths, sql)

    return RegisteredTool(
        spec=ToolSpec(
            name="query_data",
            description=TOOL_INTRO + render_schema(tables, filenames),
            parameters={
                "type": "object",
                "properties": {
                    "sql": {
                        "type": "string",
                        "description": "One SELECT / SUMMARIZE / DESCRIBE statement.",
                    }
                },
                "required": ["sql"],
            },
        ),
        handler=handler,
    )
